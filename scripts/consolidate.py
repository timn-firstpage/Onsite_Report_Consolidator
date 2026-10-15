"""CLI for explicitly selected local onsite report runs. No Google API calls."""
import argparse
import io
import json
from pathlib import Path, PurePosixPath
import stat
import sys
import zipfile

import openpyxl

from workbook import build, verify


def inventory(path):
    path=Path(path)
    def workbook_info(name,data):
        w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        info={'file':name,'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column,'headers':[c.value for c in next(s.iter_rows(max_row=1))]} for s in w]}
        w.close()
        return info
    if path.suffix.lower()=='.zip':
        with zipfile.ZipFile(path) as z:
            infos=safe_zip_infos(z)
            return [workbook_info(i.filename,z.read(i)) for i in infos if i.filename.lower().endswith('.xlsx')]
    if path.is_dir():
        files=sorted(p for p in path.iterdir() if p.suffix.lower()=='.xlsx' and not p.name.startswith(('~$','Autocheck_onsite_audit_')))
    else:files=[path]
    return [workbook_info(p.name,p.read_bytes()) for p in files]


def safe_zip_infos(z):
    infos=z.infolist()
    if len(infos)>1000 or sum(i.file_size for i in infos)>200*1024*1024:
        raise ValueError('ZIP exceeds 1000 entries / 200 MiB uncompressed limit')
    seen=set()
    for i in infos:
        name=i.filename
        p=PurePosixPath(name)
        if p.is_absolute() or '..' in p.parts or '\\' in name or ':' in name or stat.S_ISLNK(i.external_attr>>16):
            raise ValueError(f'unsafe ZIP entry: {name}')
        folded=str(p).casefold().rstrip('/')
        if folded in seen:raise ValueError(f'duplicate ZIP path: {name}')
        seen.add(folded)
    return infos


def extract(archive,destination):
    destination=Path(destination).resolve()
    if destination.exists():raise ValueError('ZIP destination already exists; choose a new run workspace')
    with zipfile.ZipFile(archive) as z:
        infos=safe_zip_infos(z)
        destination.mkdir(parents=True)
        for i in infos:
            p=destination.joinpath(*PurePosixPath(i.filename).parts)
            if not p.resolve().is_relative_to(destination):raise ValueError('ZIP path escaped destination')
            if i.is_dir():p.mkdir(parents=True,exist_ok=True);continue
            p.parent.mkdir(parents=True,exist_ok=True)
            with p.open('xb') as f:f.write(z.read(i))
    return str(destination)


def main():
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    if hasattr(sys.stderr,'reconfigure'):sys.stderr.reconfigure(encoding='utf-8')
    parser=argparse.ArgumentParser(description=__doc__)
    sub=parser.add_subparsers(dest='command',required=True)
    p=sub.add_parser('inspect',help='read workbook headers; does not select reports or build')
    p.add_argument('path')
    p=sub.add_parser('extract',help='extract a selected ZIP into a new local run workspace')
    p.add_argument('archive');p.add_argument('--destination',required=True)
    p=sub.add_parser('build',help='validate manifest inputs and build without overwrite')
    p.add_argument('--manifest',required=True)
    p=sub.add_parser('verify',help='verify output hash and internal navigation against its validation sidecar')
    p.add_argument('path')
    args=parser.parse_args()
    try:
        if args.command=='inspect':result=inventory(args.path)
        elif args.command=='extract':result={'extracted_to':extract(args.archive,args.destination)}
        elif args.command=='build':result={'output':str(build(args.manifest))}
        else:result=verify(args.path)
        print(json.dumps(result,ensure_ascii=False,indent=2,default=str))
    except (ValueError,OSError,KeyError,zipfile.BadZipFile) as e:
        print(json.dumps({'error':str(e)},ensure_ascii=False),file=sys.stderr)
        return 2
    return 0


if __name__=='__main__':raise SystemExit(main())
