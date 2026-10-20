"""Controlled template writes and offline structural verification."""
from collections import defaultdict
from copy import copy
from datetime import datetime, timezone
import io
import json
import os
from pathlib import Path
import re
import tempfile
import math
import unicodedata
from urllib.parse import urlsplit

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.workbook.views import BookView

from inputs import canonical, check_id, norm, read_run, sha
from paths import local_path
from publication import publish_no_replace
from runtime import require_runtime


def link_signature(c):
    h = c.hyperlink
    return (h.target,h.location,h.display,h.tooltip) if h else None


def cell_signature(c):
    # XLSX saves can deduplicate equivalent style-table indices; compare effective styles.
    alignment=copy(c.alignment)
    for flag in ['wrapText','shrinkToFit','justifyLastLine']:
        setattr(alignment,flag,bool(getattr(alignment,flag)))
    styles=tuple(str(v) for v in (c.font,c.fill,c.border,alignment,c.protection,c.number_format))
    value=format(c.value,'.14g') if isinstance(c.value,(int,float)) and not isinstance(c.value,bool) else c.value
    return (value,c.data_type,styles,link_signature(c),str(c.comment))


def baseline(w):
    return {s.title:{'cells':{c.coordinate:cell_signature(c) for row in s for c in row if c.value is not None or c.has_style or c.hyperlink},'merged':set(map(str,s.merged_cells.ranges)),'validation':str(s.data_validations),'images':len(s._images),'charts':len(s._charts)} for s in w}


def diagnostics(w):
    issues=[]
    for s in w:
        for row in s:
            for c in row:
                if c.data_type=='f':
                    if '#REF!' in c.value:
                        issues.append(f'{s.title}!{c.coordinate}: inherited #REF! formula')
                    for name in re.findall(r"'((?:[^']|'')+)'!",c.value):
                        if name.replace("''", "'") not in w.sheetnames:
                            issues.append(f'{s.title}!{c.coordinate}: formula references missing sheet {name}')
                    if '__xludf.DUMMYFUNCTION' in c.value:
                        issues.append(f'{s.title}!{c.coordinate}: Google-export placeholder formula, native recalculation unverified')
                if c.hyperlink and c.hyperlink.location and c.hyperlink.location.startswith(('heading=','gid=')):
                    issues.append(f'{s.title}!{c.coordinate}: inherited Google anchor {c.hyperlink.location}, not an Excel cell target')
    return issues


def internal_target(location):
    location=location.lstrip('#')
    if location.startswith(('heading=','gid=')):
        return None
    m=re.fullmatch(r"'((?:[^']|'')+)'!\$?([A-Z]+)\$?([1-9]\d*)", location)
    if not m:
        m=re.fullmatch(r'([^!]+)!\$?([A-Z]+)\$?([1-9]\d*)',location)
    if not m:
        raise ValueError(f'invalid internal link target: {location}')
    name=m[1].replace("''", "'")
    cell=m[2]+m[3]
    if openpyxl.utils.column_index_from_string(m[2])>16384 or int(m[3])>1048576:
        raise ValueError(f'out of bounds link target: {location}')
    return name,cell


def validate_links(w):
    count=0
    for s in w:
        for row in s:
            for c in row:
                locations=[]
                if c.hyperlink:
                    if c.hyperlink.location:
                        locations.append(c.hyperlink.location)
                    elif c.hyperlink.target and c.hyperlink.target.startswith('#'):
                        locations.append(c.hyperlink.target)
                if c.data_type=='f':
                    m=re.match(r'=HYPERLINK\("(#[^"]+)"',c.value,re.I)
                    if m:locations.append(m[1])
                for location in locations:
                    target=internal_target(location)
                    if target:
                        name,cell=target
                        if name not in w.sheetnames:
                            raise ValueError(f'internal link references missing sheet: {s.title}!{c.coordinate} -> {name}')
                        count+=1
    return count


def put(ws,row,col,value,touched,url=False):
    c=ws.cell(row,col)
    if c.data_type=='f':
        raise ValueError(f'refusing to overwrite formula: {ws.title}!{c.coordinate}')
    if c.hyperlink:
        raise ValueError(f'refusing to overwrite hyperlink data cell: {ws.title}!{c.coordinate}')
    c.value=value
    # Untrusted strings beginning '=' are literal report text, never formulas.
    if isinstance(value,str):c.data_type='s'
    touched.add((ws.title,c.coordinate))
    if url and isinstance(value,str) and value.strip().startswith(('http://','https://')):
        c.value=value.strip()
        c.hyperlink=c.value
        c.font=Font(name=c.font.name or 'Arial',size=c.font.sz or 10,color='0563C1',underline='single')
    return c


def navigation(c,sheet,cell):
    name=sheet.replace("'","''")
    c.hyperlink=Hyperlink(ref=c.coordinate,location=f"'{name}'!{cell}")


def validate_navigation(w,links):
    for name,coord,sheet,cell in links:
        source=w[name][coord]
        if (not source.hyperlink or source.hyperlink.target or
                internal_target(source.hyperlink.location or '')!=(sheet,cell)):
            raise ValueError(f'navigation target mismatch: {name}!{coord}')
        if sheet not in w.sheetnames or w[sheet][cell].value is None:
            raise ValueError(f'navigation target missing or empty: {sheet}!{cell}')


def evidence_label(check):
    n=check['detail_count']
    if check['status'] in {'-','NA','Human check'}:
        return f'问题数量未确认；已提供明细 {n} 条'
    if check['status']=='✔':
        return '来源检查通过；已记录问题明细 0 条（不代表独立复核）'
    return f'已记录问题明细 {n} 条（不代表唯一问题总数）'


def standalone_url(value):
    return isinstance(value,str) and re.fullmatch(r'https?://\S+',value.strip()) and bool(urlsplit(value.strip()).hostname)


def fit_row(ws,rownum,minimum=36):
    height=minimum
    for cell in ws[rownum]:
        if cell.value is None:continue
        width=ws.sheet_format.defaultColWidth or 13
        for dimension in ws.column_dimensions.values():
            lo=dimension.min or openpyxl.utils.column_index_from_string(dimension.index)
            hi=dimension.max or lo
            if lo<=cell.column<=hi:width=dimension.width;break
        lines=0
        for line in str(cell.value).split('\n'):
            length=sum(2 if unicodedata.east_asian_width(ch) in {'W','F'} else 1 for ch in line)
            lines+=max(1,math.ceil(length/max(1,width-2)))
        height=max(height,lines*(cell.font.sz or 10)*1.35+6)
    ws.row_dimensions[rownum].height=min(409,height)


def target_rows(main):
    rows={}
    for row in main.iter_rows():
        value=row[1].value if len(row)>1 else None
        if value is None:continue
        try:ident=check_id(value)
        except ValueError:continue
        if ident in rows:raise ValueError(f'duplicate template check {ident}')
        rows[ident]=row[0].row
    return rows


def extend_human_validation(main,rows):
    """Extend onsite status lists while preserving excluded GSC cells' original list."""
    gsc_rows={r for ident,r in rows.items() if ident.startswith('7.')}
    for validation in list(main.data_validations.dataValidation):
        formula=str(validation.formula1 or '')
        if validation.type!='list' or not formula.startswith('"') or not all(x in formula for x in ['✔','✖','NA','-']):continue
        tokens=formula.strip('"').split(',')
        if 'Human check' in tokens:continue
        keep,extend=[],[]
        for area in validation.sqref.ranges:
            if not area.min_col<=4<=area.max_col:
                keep.append(str(area));continue
            if area.min_col<4:keep.append(str(CellRange(min_col=area.min_col,max_col=3,min_row=area.min_row,max_row=area.max_row)))
            if area.max_col>4:keep.append(str(CellRange(min_col=5,max_col=area.max_col,min_row=area.min_row,max_row=area.max_row)))
            start=area.min_row
            for r in sorted(r for r in gsc_rows if area.min_row<=r<=area.max_row):
                if start<r:extend.append(f'D{start}:D{r-1}')
                keep.append(f'D{r}');start=r+1
            if start<=area.max_row:extend.append(f'D{start}:D{area.max_row}')
        if not extend:continue
        updated=copy(validation);updated.sqref=' '.join(extend)
        updated.formula1='"'+','.join(tokens+['Human check'])+'"'
        if keep:validation.sqref=' '.join(keep)
        else:main.data_validations.dataValidation.remove(validation)
        main.add_data_validation(updated)


def rename_main(w,old,new):
    if old==new:return
    if not new or len(new)>31 or re.search(r'[\\/*?:\[\]]',new):
        raise ValueError('output_main_sheet must be a valid Excel worksheet name of at most 31 characters')
    if any(n.casefold()==new.casefold() for n in w.sheetnames):
        raise ValueError(f'output main worksheet name already exists: {new}')
    old_ref="'"+old.replace("'","''")+"'!"
    new_ref="'"+new.replace("'","''")+"'!"
    def rewrite(value):
        return value.replace(old_ref,new_ref) if isinstance(value,str) else value
    w[old].title=new
    for s in w:
        for row in s:
            for c in row:
                if c.data_type=='f':c.value=rewrite(c.value)
                if c.hyperlink:
                    c.hyperlink.location=rewrite(c.hyperlink.location)
                    if c.hyperlink.target and c.hyperlink.target.startswith('#'):
                        c.hyperlink.target=rewrite(c.hyperlink.target)
        for defined in s.defined_names.values():defined.attr_text=rewrite(defined.attr_text)
        for validation in s.data_validations.dataValidation:
            validation.formula1=rewrite(validation.formula1)
            validation.formula2=rewrite(validation.formula2)
        for rule_set in s.conditional_formatting:
            for rule in s.conditional_formatting[rule_set]:
                if rule.formula:rule.formula=[rewrite(f) for f in rule.formula]
        for chart in s._charts:
            for series in chart.series:
                for field in ['val','cat','xVal','yVal','zVal','tx']:
                    data=getattr(series,field,None)
                    for ref in ['numRef','strRef','multiLvlStrRef']:
                        obj=getattr(data,ref,None)
                        if obj is not None and getattr(obj,'f',None):obj.f=rewrite(obj.f)
    for defined in w.defined_names.values():defined.attr_text=rewrite(defined.attr_text)


def write_details(w,run,touched):
    main_name=run['profile']['main_sheet']
    main_rows=target_rows(w[main_name])
    grouped=defaultdict(list)
    for ident,c in run['checks'].items():
        topic=run['profile']['topics'][c['topic']]
        grouped[topic['targets'][ident]].append(c)
    detail_locations={}
    detail_records={}
    evidence_cells=[]
    for name,checks in grouped.items():
        if name not in w.sheetnames:
            raise ValueError(f'missing destination worksheet; update explicit profile: {name}')
        ws=w[name]
        header_row=4 if name=='10. Robots.txt' else 3
        headers={c.column:str(c.value) for c in ws[header_row] if c.value is not None}
        if not headers:raise ValueError(f'missing destination headers in {name}')
        all_fields=[]
        for c in checks:
            for d in c['details']:
                for field in d['values']:
                    if field not in all_fields:all_fields.append(field)
        mapping={}
        for col,h in headers.items():
            choices=[f for f in all_fields if norm(f)==norm(h)]
            if not choices:
                choices=[f for f in all_fields if canonical(f,run['profile'])==canonical(h,run['profile']) and not canonical(h,run['profile']).startswith('extra:')]
            if choices:
                mapping[col]=choices[0]
        occupied=[c.row for row in ws.iter_rows(min_row=header_row+1) for c in row if c.value is not None]
        start=max(occupied)+2 if occupied else header_row+1
        rownum=start
        for check in checks:
            if check['id'] not in main_rows:
                raise ValueError(f'checks missing from template: {check["id"]}')
            detail_locations[check['id']]=(name,ws.cell(header_row,min(headers)).coordinate)
            detail_records[check['id']]=[]
            for detail in check['details']:
                output_cells=[]
                for col,field in mapping.items():
                    value=detail['values'].get(field)
                    is_url=canonical(field,run['profile']) in {'address','destination','sitemap'} or standalone_url(value)
                    c=put(ws,rownum,col,value,touched,url=is_url);c.alignment=Alignment(vertical='top',wrap_text=True)
                    if value is not None:output_cells.append(c.coordinate)
                    if is_url and value is not None:evidence_cells.append([name,c.coordinate,str(value).strip()])
                if not output_cells:
                    raise ValueError(f'no source detail fields map to existing template schema: {name}, check {check["id"]}')
                if not detail_records[check['id']]:detail_locations[check['id']]=(name,output_cells[0])
                detail_records[check['id']].append({'source_file':check['source_file'],'source_check':check['id'],'source_sheet':detail['source_sheet'],'source_row':detail['source_row'],'source_values':detail['values'],'output_sheet':name,'output_row':rownum,'output_cells':output_cells})
                fit_row(ws,rownum)
                rownum+=1
        # Rebuild controlled navigation even if a template link/formula already exists.
        if ws['A1'].data_type=='f' or ws['A1'].value is None:
            ws['A1']='Implementation Checklist'
        touched.add((name,'A1'))
        navigation(ws['A1'],main_name,f'B{min(main_rows[c["id"]] for c in checks)}')
    return detail_locations,evidence_cells,detail_records


def validate_output(path,original,touched,expected,evidence_cells,template_warnings,write_manifest,navigation_links):
    w=openpyxl.load_workbook(path,data_only=False)
    if w.sheetnames!=list(original):raise ValueError('template worksheet order/names changed')
    if expected and w.active.title!=expected[0][0]:raise ValueError('output does not open on main Checklist')
    for name,b in original.items():
        s=w[name]
        if set(map(str,s.merged_cells.ranges))!=b['merged'] or str(s.data_validations)!=b['validation']:
            raise ValueError(f'template merges/validation changed: {name}')
        if len(s._images)!=b['images'] or len(s._charts)!=b['charts']:
            raise ValueError(f'template images/charts changed: {name}')
        for coord,signature in b['cells'].items():
            if (name,coord) not in touched and cell_signature(s[coord])!=signature:
                raise ValueError(f'unexpected template change: {name}!{coord}')
    for name,coord,value in expected:
        if (w[name][coord].value or None)!=(value or None):raise ValueError(f'output value mismatch {name}!{coord}')
    for name,coord,value,links in write_manifest:
        c=w[name][coord]
        if (c.value or None)!=(value or None) or link_signature(c)!=links:
            raise ValueError(f'written detail/navigation did not survive export: {name}!{coord}')
    for name,coord,url in evidence_cells:
        c=w[name][coord]
        if c.value!=url or not c.hyperlink or c.hyperlink.target!=url:
            raise ValueError(f'website URL hyperlink missing: {name}!{coord}')
    new_warnings=set(diagnostics(w))-set(template_warnings)
    if new_warnings:raise ValueError(f'new formula/link errors: {sorted(new_warnings)}')
    count=validate_links(w)
    validate_navigation(w,navigation_links)
    w.close()
    return count


def build(manifest_path):
    run=read_run(manifest_path)
    m=run['manifest']
    output=run['folder']/f'Autocheck_onsite_audit_{m["company"]}_{m["output_date"]}.xlsx'
    sidecar=output.with_suffix('.validation.json')
    if output.exists() or sidecar.exists():raise ValueError(f'output already exists: {output}')
    w=openpyxl.load_workbook(io.BytesIO(run['template'].read_bytes()),data_only=False)
    main_name=run['profile']['main_sheet']
    if main_name not in w.sheetnames:raise ValueError(f'missing main worksheet {main_name}')
    output_main=run['profile'].get('output_main_sheet',main_name)
    rename_main(w,main_name,output_main)
    main_name=output_main
    run['profile']['main_sheet']=main_name
    main=w[main_name]
    rows=target_rows(main)
    extend_human_validation(main,rows)
    original=baseline(w)
    warnings=diagnostics(w)
    # Validate after rebuilding controlled links; stale template navigation may be broken.
    touched=set()
    locations,evidence_cells,detail_records=write_details(w,run,touched)
    expected=[]
    missing=[]
    mapped_ids=set(i for topic in run['profile']['topics'].values() for i in topic['ids'])
    status_cells=set()
    for validation in main.data_validations.dataValidation:
        if validation.type=='list' and all(x in str(validation.formula1) for x in ['✔','✖','NA','-']):
            status_cells.update(f'D{r}' for r in rows.values() if f'D{r}' in validation.sqref)
    for ident in rows:
        if ident.startswith('7.'):continue
        if ident in mapped_ids or f'D{rows[ident]}' in status_cells:
            if ident not in rows:continue
            r=rows[ident]
            check=run['checks'].get(ident)
            if check:
                findings=check['findings']+'\n[汇总] '+evidence_label(check)
                if check['status']=='-':findings+='；来源状态：'+check['raw_status']
                values=[check['status'],findings,check['instruction']]
            else:
                values=['-','未检查：本次未提供该检查项报告。','']
                missing.append(ident)
            for col,value in zip([4,5,6],values):
                c=put(main,r,col,value,touched)
                if col in {5,6}:c.alignment=copy(c.alignment);c.alignment=Alignment(horizontal=c.alignment.horizontal,vertical='top',wrap_text=True)
                expected.append([main_name,c.coordinate,value])
            if check:
                sheet,cell=locations[ident]
                c=main.cell(r,7)
                if c.data_type=='f':c.value=sheet
                if c.value is None:put(main,r,7,sheet,touched)
                touched.add((main_name,c.coordinate));navigation(c,sheet,cell)
            fit_row(main,r,max(60,main.row_dimensions[r].height or 0))
    absent=set(run['checks'])-set(rows)
    if absent:raise ValueError(f'checks missing from template: {sorted(absent)}')
    w.active=w.sheetnames.index(main_name)
    if not w.views:w.views=[BookView(activeTab=w.sheetnames.index(main_name))]
    fd,tmp=tempfile.mkstemp(prefix='.consolidating-',suffix='.xlsx',dir=run['folder']);os.close(fd)
    temp_path=Path(tmp)
    write_manifest=[[name,coord,w[name][coord].value,link_signature(w[name][coord])] for name,coord in sorted(touched)]
    navigation_links=[]
    for name,coord in sorted(touched):
        h=w[name][coord].hyperlink
        if h and h.location and not h.target:
            target=internal_target(h.location)
            if target:navigation_links.append([name,coord,*target])
    try:
        w.save(temp_path);w.close()
        link_count=validate_output(temp_path,original,touched,expected,evidence_cells,warnings,write_manifest,navigation_links)
        for source in run['sources']:
            if sha(Path(source['path']).read_bytes())!=source['sha256']:raise ValueError('source changed during build')
        for p,digest in [(run['template'],run['template_sha256']),(Path(run['manifest_path']),run['manifest_sha256']),(Path(run['profile_path']),run['profile_sha256'])]:
            if sha(p.read_bytes())!=digest:raise ValueError(f'configuration changed during build: {p}')
        report={'schema_version':1,'run_id':m['run_id'],'company':m['company'],'site_hosts':m['site_hosts'],'output_date':m['output_date'],'created_utc':datetime.now(timezone.utc).isoformat(),'output_sha256':sha(temp_path.read_bytes()),'template_sha256':run['template_sha256'],'manifest_sha256':run['manifest_sha256'],'profile_sha256':run['profile_sha256'],'sources':run['sources'],'missing_checks':missing,'checks':{i:{'status':c['status'],'raw_status':c['raw_status'],'recorded_detail_rows':c['detail_count'],'source_file':c['source_file'],'detail_target':locations[i]} for i,c in run['checks'].items()},'validated_internal_links':link_count,'validated_website_links':len(evidence_cells),'template_warnings':warnings,'input_warnings':run['warnings'],'native_excel_recalculation':'not performed'}
        report['publication_method']=publish_no_replace(temp_path,output)
        report['navigation_links']=navigation_links
        for ident,check in run['checks'].items():
            report['checks'][ident].update({'detail_target':locations[ident],'details':detail_records[ident],'source_sheet':check['source_sheet'],'source_row':check['source_row'],'source_date':check['source_date'],'source_check_name':check['source_check_name'],'findings':check['findings'],'instruction':check['instruction'],'extra_checklist':check['extra_checklist']})
        try:
            with sidecar.open('x',encoding='utf-8') as f:json.dump(report,f,ensure_ascii=False,indent=2)
        except Exception:
            output.unlink();raise
    finally:
        temp_path.unlink(missing_ok=True)
    return output


def verify(path):
    require_runtime()
    path=local_path(path)
    report=json.loads(path.with_suffix('.validation.json').read_text(encoding='utf-8'))
    if sha(path.read_bytes())!=report['output_sha256']:raise ValueError('output hash changed since validation')
    w=openpyxl.load_workbook(path,data_only=False)
    links=validate_links(w)
    validate_navigation(w,report.get('navigation_links',[]))
    w.close()
    return {'output':str(path),'sha256':report['output_sha256'],'checks':len(report['checks']),'internal_links':links,'website_links':report['validated_website_links'],'template_warning_count':len(report['template_warnings'])}
