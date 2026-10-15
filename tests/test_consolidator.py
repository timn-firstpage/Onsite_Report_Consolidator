import copy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_MAIN = '(auto) SEO Checklist'
sys.path.insert(0, str(ROOT / 'scripts'))


class ConsolidationTests(unittest.TestCase):
    def setUp(self):
        (ROOT / '.runs').mkdir(exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=ROOT / '.runs')
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.template = self.folder / 'template.xlsx'
        w = openpyxl.Workbook()
        s = w.active
        s.title = 'SEO Implementation Checklist'
        s.append(['', '#', 'ITEM', 'Initial Check', 'Findings', 'Instruction', 'Links'])
        for ident in ['7.1', '9.1', '9.2', '9.3'] + [f'11.{i}' for i in range(1, 10)] + [f'17.{i}' for i in range(1, 5)]:
            s.append(['', ident, ident + ' question'])
        s['D2'] = 'original GSC'
        s['J2'] = '=1+2'
        s['J2'].fill = PatternFill('solid', fgColor='123456')
        for title, headers in [('11. Sitemap', ['Sitemaps', 'URLs Example', 'Instructions', 'Screenshots (If Applicable)']), ('9. Implement HTTPS Protocol', ['Source','Destination','Anchor Text','Issues','Screenshots (If have)','Specific Instructions (If have)'])] + [(title, ['Address','Titles','New Title','Instructions']) for title in ['17.1 Page titles - Missing','17.2 Page titles - Duplicate','17.3 Page Titles - Too Long','17.4 Page Titles - Multiple']]:
            s = w.create_sheet(title)
            s['A1'] = 'Implementation Checklist'
            s['A1'].hyperlink = "#'SEO Implementation Checklist'!A1"
            s.append(['Priorities & Status'])
            s.append(headers)
        w['11. Sitemap']['C4'] = 'preserve template note'
        w.save(self.template)
        self.manifest = {'run_id':'test-run', 'run_folder':'.', 'company':'Example Co', 'site_hosts':['www.example.com'], 'output_date':'2026-10-15', 'template':'template.xlsx', 'sources':[]}
        self.path = self.folder / 'run.json'

    def report(self, topic='sitemap_audit', ids=None, detail=None, statuses=None, split=False, legacy=False):
        ids = ids or [f'11.{i}' for i in range(1,10)]
        name = ('Example Co_' if legacy else f'{ids[0]}-{ids[-1]}_Example Co_') + topic + '_2026-10-13.xlsx'
        p = self.folder / name
        w = openpyxl.Workbook()
        s = w.active
        s.title = 'Checklist'
        s.append(['Item No.','Item Name','Initial Check','Findings','Instruction'])
        for ident in ids:
            s.append([ident,ident+' question',(statuses or {}).get(ident,'√'),'Source observation','Source instruction'])
        if topic == 'sitemap_audit':
            s = w.create_sheet('11. Sitemap')
            s.append(['Sitemaps','URLs Example','Instructions','Screenshots (If Applicable)'])
        else:
            if split:
                for ident in ids:
                    s=w.create_sheet(ident+' detail')
                    s.append(['Address','Title 1','New Title','Instruction'])
                    for row in detail or []:
                        if row[0]==ident:s.append(row[1:])
            else:
                s=w.create_sheet('17. Page Title')
                s.append(['Item No.','Address','Title 1','New Title','Instruction'])
                for row in detail or []:s.append(row)
        w.save(p)
        self.manifest['sources'].append({'file':name,'topic':topic,'date':'2026-10-13','legacy_name':legacy})
        return p

    def save(self):
        self.path.write_text(json.dumps(self.manifest),encoding='utf-8')
        return self.path

    def read(self):
        from inputs import read_run
        return read_run(self.save())

    def build(self):
        from workbook import build
        return build(self.save())

    def test_human_status_and_original_findings_survive(self):
        self.report(statuses={'11.3':'Human check'})
        r=self.read()
        self.assertEqual(r['checks']['11.3']['status'],'Human check')
        self.assertEqual(r['checks']['11.3']['raw_status'],'Human check')
        self.assertEqual(r['checks']['11.3']['findings'],'Source observation')

    def test_unknown_status_stops(self):
        self.report(statuses={'11.3':'maybe'})
        with self.assertRaisesRegex(ValueError,'status'):self.read()

    def test_legacy_filename_requires_explicit_permission(self):
        self.report(legacy=True)
        self.manifest['sources'][0]['legacy_name']=False
        with self.assertRaisesRegex(ValueError,'filename'):self.read()

    def test_wrong_filename_company_stops(self):
        self.report()
        self.manifest['company']='Other Co'
        with self.assertRaisesRegex(ValueError,'company'):self.read()

    def test_date_mismatch_stops(self):
        self.report()
        self.manifest['sources'][0]['date']='2026-10-14'
        with self.assertRaisesRegex(ValueError,'date'):self.read()

    def test_duplicate_candidates_stop(self):
        self.report()
        self.manifest['sources'].append(copy.deepcopy(self.manifest['sources'][0]))
        with self.assertRaisesRegex(ValueError,'duplicate'):self.read()

    def test_missing_check_in_full_report_stops(self):
        self.report(ids=['11.1'])
        with self.assertRaisesRegex(ValueError,'coverage|missing'):self.read()

    def test_explicit_subset_report_is_supported(self):
        self.report(ids=['11.1'])
        self.manifest['sources'][0]['checks']=['11.1']
        self.assertEqual(list(self.read()['checks']),['11.1'])

    def test_unknown_headers_stop(self):
        p=self.report()
        w=openpyxl.load_workbook(p);w['Checklist']['C1']='Unrecognized';w.save(p)
        with self.assertRaisesRegex(ValueError,'column|header'):self.read()

    def test_wrong_site_in_detail_stops(self):
        self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://other.example/title','old','new','fix']], {'17.2':'X'})
        with self.assertRaisesRegex(ValueError,'host'):self.read()

    def test_gsc_explicitly_selected_is_rejected(self):
        self.manifest['sources']=[{'file':'gsc.xlsx','topic':'gsc','date':'2026-10-13'}]
        with self.assertRaisesRegex(ValueError,'excluded|GSC|topic'):self.read()

    def test_combined_details_route_to_exact_subitem(self):
        self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        r=self.read()
        self.assertEqual(len(r['checks']['17.2']['details']),1)
        self.assertEqual(len(r['checks']['17.1']['details']),0)

    def test_split_details_work_with_reordered_columns(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'},split=True)
        w=openpyxl.load_workbook(p);s=w['17.2 detail'];s['E1']='Extra evidence';s['E2']='keep me';w.save(p)
        self.assertEqual(self.read()['checks']['17.2']['details'][0]['values']['Extra evidence'],'keep me')

    def test_source_count_disagreement_stops(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        w=openpyxl.load_workbook(p);w['Checklist']['D3']='2 issue row(s) found.';w.save(p)
        with self.assertRaisesRegex(ValueError,'count'):self.read()

    def test_pass_with_problem_details_stops(self):
        self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']])
        with self.assertRaisesRegex(ValueError,'result|status'):self.read()

    def test_failed_result_without_any_detail_stops(self):
        self.report(statuses={'11.3':'X'})
        with self.assertRaisesRegex(ValueError,'detail'):self.read()

    def test_unlisted_files_are_not_read(self):
        self.report()
        (self.folder/'unrelated.xlsx').write_bytes(b'invalid and not selected')
        self.assertEqual(len(self.read()['checks']),9)

    def test_path_escape_stops(self):
        self.report();self.manifest['sources'][0]['file']='../outside.xlsx'
        with self.assertRaisesRegex(ValueError,'outside|relative'):self.read()

    def test_build_preserves_template_and_gsc(self):
        self.report(statuses={'11.3':'Human check'})
        p=self.build();w=openpyxl.load_workbook(p)
        self.assertEqual(p.name,'Autocheck_onsite_audit_Example Co_2026-10-15.xlsx')
        s=w[OUTPUT_MAIN]
        self.assertEqual(s['D2'].value,'original GSC')
        self.assertEqual(s['J2'].value,'=1+2')
        self.assertEqual(s['J2'].fill.fgColor.rgb,'00123456')
        self.assertEqual(s['D3'].value,'-')
        self.assertIn('未检查',s['E3'].value)
        self.assertEqual(w['11. Sitemap']['C4'].value,'preserve template note')
        self.assertEqual(openpyxl.load_workbook(self.template)['SEO Implementation Checklist']['D6'].value,None)

    def test_count_and_hyperlink_and_source_traceability(self):
        self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        p=self.build();w=openpyxl.load_workbook(p)
        s=w['17.2 Page titles - Duplicate']
        self.assertEqual(s['A4'].value,'https://www.example.com/a')
        self.assertEqual(s['A4'].hyperlink.target,'https://www.example.com/a')
        self.assertTrue(any(c.value=='17.2' for row in s for c in row))
        self.assertTrue(any('page_titles_audit' in str(c.value) for row in s for c in row))
        main=w[OUTPUT_MAIN]
        r=next(row[0].row for row in main if row[1].value=='17.2')
        self.assertEqual(main.cell(r,4).value,'✖')
        self.assertIn('1',main.cell(r,5).value)
        self.assertIn('17.2 Page titles',main.cell(r,7).hyperlink.location)

    def test_no_output_overwrite(self):
        self.report();p=self.build();original=p.read_bytes()
        with self.assertRaisesRegex(ValueError,'exists'):self.build()
        self.assertEqual(original,p.read_bytes())

    def test_output_tampering_detected(self):
        from workbook import verify
        self.report();p=self.build()
        w=openpyxl.load_workbook(p);w[OUTPUT_MAIN]['D6']='✖';w.save(p)
        with self.assertRaisesRegex(ValueError,'hash|changed'):verify(p)

    def test_asset_host_does_not_authorize_foreign_page_urls(self):
        self.manifest['asset_hosts']=['other.example']
        self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://other.example/title','old','new','fix']], {'17.2':'X'})
        with self.assertRaisesRegex(ValueError,'host'):self.read()

    def test_split_sheet_and_internal_id_conflict_stops(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'},split=True)
        w=openpyxl.load_workbook(p);s=w['17.2 detail'];s['F1']='Item No.';s['F2']='17.1';w.save(p)
        with self.assertRaisesRegex(ValueError,'mismatch'):self.read()

    def test_template_formula_in_data_area_is_not_overwritten(self):
        self.report()
        w=openpyxl.load_workbook(self.template);w['SEO Implementation Checklist']['E6']='=1+1';w.save(self.template)
        with self.assertRaisesRegex(ValueError,'formula'):self.build()

    def test_template_hyperlink_in_data_area_is_not_overwritten(self):
        self.report()
        w=openpyxl.load_workbook(self.template);s=w['SEO Implementation Checklist'];s['E6']='Keep link';s['E6'].hyperlink='https://www.example.com/';w.save(self.template)
        with self.assertRaisesRegex(ValueError,'hyperlink'):self.build()

    def test_source_formula_is_not_executed(self):
        p=self.report()
        w=openpyxl.load_workbook(p);w['Checklist']['D2']='=HYPERLINK("https://other.example", "run")';w.save(p)
        with self.assertRaisesRegex(ValueError,'formula'):self.read()

    def test_longer_topic_alias_is_not_misread_as_company(self):
        p=self.report()
        renamed=p.with_name(p.name.replace('_sitemap_audit_', '_onsite_sitemap_audit_'))
        p.rename(renamed)
        self.manifest['sources'][0]['file']=renamed.name
        profile=json.loads((ROOT/'scripts/profiles.json').read_text())
        profile['topics']['sitemap_audit']['aliases'].append('onsite_sitemap_audit')
        (self.folder/'profile.json').write_text(json.dumps(profile))
        self.manifest['profile']='profile.json'
        self.assertEqual(len(self.read()['checks']),9)

    def test_populated_detail_cell_without_header_stops(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        w=openpyxl.load_workbook(p);w['17. Page Title']['G2']='material evidence';w.save(p)
        with self.assertRaisesRegex(ValueError,'header'):self.read()

    def test_plain_plural_source_count_disagreement_stops(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        w=openpyxl.load_workbook(p);w['Checklist']['D3']='2 issue rows found.';w.save(p)
        with self.assertRaisesRegex(ValueError,'count'):self.read()

    def test_merged_source_detail_requires_adapter(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        w=openpyxl.load_workbook(p);w['17. Page Title'].merge_cells('C2:D2');w.save(p)
        with self.assertRaisesRegex(ValueError,'merged'):self.read()

    def test_merged_source_checklist_requires_adapter(self):
        p=self.report()
        w=openpyxl.load_workbook(p);w['Checklist'].merge_cells('D2:E2');w.save(p)
        with self.assertRaisesRegex(ValueError,'merged'):self.read()

    def test_incomplete_evidence_keeps_explicit_failed_result(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        w=openpyxl.load_workbook(p);w['Checklist']['D3']='1 issue row(s) found. Evidence incomplete.';w.save(p)
        r=self.read()
        self.assertEqual(r['checks']['17.2']['status'],'✖')
        self.assertIn('Evidence incomplete',r['checks']['17.2']['findings'])

    def test_declared_asset_destination_is_clickable(self):
        self.manifest['asset_hosts']=['cdn.example.com']
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        w=openpyxl.load_workbook(p);s=w['17. Page Title'];s['F1']='Destination';s['F2']='https://cdn.example.com/file.png';w.save(p)
        out=self.build();w=openpyxl.load_workbook(out)
        cells=[c for row in w['17.2 Page titles - Duplicate'] for c in row if c.value=='https://cdn.example.com/file.png']
        self.assertEqual(len(cells),1)
        self.assertEqual(cells[0].hyperlink.target,'https://cdn.example.com/file.png')

    def test_extra_evidence_url_is_clickable(self):
        p=self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a','old','new','fix']], {'17.2':'X'})
        w=openpyxl.load_workbook(p);s=w['17. Page Title'];s['F1']='Evidence URL';s['F2']='https://evidence.example/screenshot';w.save(p)
        out=self.build();w=openpyxl.load_workbook(out)
        c=next(c for row in w['17.2 Page titles - Duplicate'] for c in row if c.value=='https://evidence.example/screenshot')
        self.assertEqual(c.hyperlink.target,'https://evidence.example/screenshot')

    def test_missing_other_template_check_marked_unchecked(self):
        from openpyxl.worksheet.datavalidation import DataValidation
        w=openpyxl.load_workbook(self.template);s=w['SEO Implementation Checklist'];s.append(['','21.1','Other onsite check'])
        row=s.max_row;dv=DataValidation(type='list',formula1='"✔,✖,NA,-"');s.add_data_validation(dv);dv.add(s.cell(row,4));w.save(self.template)
        self.report();out=self.build();s=openpyxl.load_workbook(out)[OUTPUT_MAIN]
        self.assertEqual(s.cell(row,4).value,'-')
        self.assertIn('未检查',s.cell(row,5).value)

    def test_named_extra_checklist_field_survives(self):
        p=self.report()
        w=openpyxl.load_workbook(p);s=w['Checklist'];s['F1']='Evidence limit';s['F2']='Sample only';w.save(p)
        out=self.build();s=openpyxl.load_workbook(out)['11. Sitemap']
        self.assertTrue(any(c.value=='Sample only' for row in s for c in row))

    def test_missing_image_destination_column_stops(self):
        w=openpyxl.Workbook();s=w.active;s.title='Checklist'
        s.append(['Item No.','Item Name','Initial Check','Findings','Instruction'])
        s.append(['20.1','Missing alt','X','1 issue row(s) found.','fix'])
        s.append(['20.2','Oversize','√','No issue found.',''])
        s=w.create_sheet('20. Image');s.append(['Item No.','Source']);s.append(['20.1','https://www.example.com/'])
        name='20.1-20.2_Example Co_images_audit_2026-10-13.xlsx';w.save(self.folder/name)
        self.manifest['sources']=[{'file':name,'topic':'images_audit','date':'2026-10-13'}]
        with self.assertRaisesRegex(ValueError,'column|field'):self.read()

    def test_unencoded_whitespace_in_website_url_stops(self):
        self.report('page_titles_audit',[f'17.{i}' for i in range(1,5)], [['17.2','https://www.example.com/a b','old','new','fix']], {'17.2':'X'})
        with self.assertRaisesRegex(ValueError,'URL'):self.read()

    def test_human_check_validation_and_missing_status_are_distinct(self):
        from openpyxl.worksheet.datavalidation import DataValidation
        w=openpyxl.load_workbook(self.template);s=w['SEO Implementation Checklist']
        dv=DataValidation(type='list',formula1='"✔,✖,NA,-"');s.add_data_validation(dv);dv.add('D2:D18');w.save(self.template)
        self.report(statuses={'11.3':'Human Check'})
        out=self.build();w=openpyxl.load_workbook(out);s=w[OUTPUT_MAIN]
        self.assertEqual(s['D8'].value,'Human check')
        self.assertEqual(s['D3'].value,'-')
        self.assertIn('未检查',s['E3'].value)
        self.assertIn('问题数量未确认',s['E8'].value)
        human_rules=[v for v in s.data_validations.dataValidation if 'D8' in v.sqref]
        self.assertEqual(len(human_rules),1)
        self.assertIn('Human check',human_rules[0].formula1)
        gsc_rules=[v for v in s.data_validations.dataValidation if 'D2' in v.sqref]
        self.assertEqual(len(gsc_rules),1)
        self.assertNotIn('Human check',gsc_rules[0].formula1)
        detail=w['11. Sitemap']
        summary=next(row for row in detail if row[0].value=='11.3')
        self.assertEqual(summary[1].value,'Human check')
        self.assertIsNone(summary[2].value)

    def test_main_rename_updates_formulas_named_ranges_and_return_links(self):
        from openpyxl.workbook.defined_name import DefinedName
        w=openpyxl.load_workbook(self.template);s=w['11. Sitemap']
        s['L2']="='SEO Implementation Checklist'!C1"
        w.defined_names.add(DefinedName('ChecklistTitle',attr_text="'SEO Implementation Checklist'!$C$1"))
        w.save(self.template)
        self.report();out=self.build();w=openpyxl.load_workbook(out)
        self.assertIn(OUTPUT_MAIN,w.sheetnames)
        self.assertNotIn('SEO Implementation Checklist',w.sheetnames)
        self.assertEqual(w['11. Sitemap']['L2'].value,"='(auto) SEO Checklist'!C1")
        self.assertEqual(w.defined_names['ChecklistTitle'].attr_text,"'(auto) SEO Checklist'!$C$1")
        self.assertIn("'(auto) SEO Checklist'!",w['11. Sitemap']['A1'].hyperlink.target)
        self.assertIn('SEO Implementation Checklist',openpyxl.load_workbook(self.template).sheetnames)

    def test_main_is_active_when_template_has_no_workbook_view(self):
        w=openpyxl.load_workbook(self.template);w.create_sheet('Summary',0);w.views=[];w.save(self.template)
        self.report();out=self.build();w=openpyxl.load_workbook(out)
        self.assertEqual(w.active.title,OUTPUT_MAIN)


if __name__=='__main__':unittest.main()
